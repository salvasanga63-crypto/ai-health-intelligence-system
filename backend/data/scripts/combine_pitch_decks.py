from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pathlib import Path

root = Path('/home/salva-kil/PROJECTS/AI-driven Health Intelligence System')
out_file = root / 'AI_Health_Intelligence_Combined_Pitch_Deck.pptx'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(12, 41, 68)
BLUE = RGBColor(0, 102, 204)
TEAL = RGBColor(0, 150, 136)
LIGHT = RGBColor(245, 248, 252)
DARK = RGBColor(35, 45, 55)
GRAY = RGBColor(100, 112, 124)
WHITE = RGBColor(255, 255, 255)


def add_title(slide, title, subtitle=None):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = LIGHT
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.0), Inches(0.8))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.05), Inches(9.0), Inches(0.5))
        sf = sub_box.text_frame
        sf.clear()
        p2 = sf.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(11)
        p2.font.color.rgb = GRAY


def add_bullets(slide, bullets, top=1.6, left=0.8, width=6.3, height=5.0):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18)
        p.font.color.rgb = DARK
        p.font.name = 'Calibri'
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(7)


def add_box(slide, x, y, w, h, title, body, fill_color):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    tx = shape.text_frame
    tx.clear()
    p = tx.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tx.add_paragraph()
    p2.text = body
    p2.font.size = Pt(11)
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(4)

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
textbox = slide.shapes.add_textbox(Inches(0.7), Inches(0.9), Inches(8.5), Inches(1.2))
tf = textbox.text_frame; tf.clear(); p = tf.paragraphs[0]; p.text = 'AI-Driven Health Intelligence System'; p.font.size = Pt(28); p.font.bold = True; p.font.color.rgb = WHITE
sub = slide.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(9.8), Inches(1.0)); sf = sub.text_frame; sf.clear(); p2 = sf.paragraphs[0]; p2.text = 'A smart healthcare platform for intelligent triage, diagnostics, patient management, and predictive maintenance.'; p2.font.size = Pt(17); p2.font.color.rgb = RGBColor(225,233,242)
shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(3.0), Inches(3.7), Inches(2.0)); shape.fill.solid(); shape.fill.fore_color.rgb = TEAL; shape.line.color.rgb = TEAL; shape.text_frame.text = 'Supervisor Evaluation Deck\nIntegrated from two presentations'
for para in shape.text_frame.paragraphs:
    para.font.size = Pt(16); para.font.color.rgb = WHITE; para.alignment = PP_ALIGN.CENTER

# Slide 2: Problem
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'Problem Statement', 'Why this platform matters')
add_bullets(slide, [
    'Healthcare systems often struggle with delayed triage and fragmented patient information.',
    'Maintenance issues in medical equipment can disrupt critical services and care quality.',
    'Manual processes make it hard to forecast demand, manage risk, and support clinicians efficiently.',
    'There is a strong need for an integrated, intelligent, scalable digital health platform.'
])

# Slide 3: Solution Overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'Proposed Solution', 'An integrated AI-powered healthcare platform')
add_bullets(slide, [
    'Combines AI diagnostics, predictive analytics, and patient management in one system.',
    'Provides role-based dashboards for administrators, doctors, nurses, and patients.',
    'Supports secure access, clinical communication, and operational visibility.',
    'Designed to improve decision support and healthcare service delivery.'
], width=7.0)

# Slide 4: Key Features
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'Key Features')
add_box(slide, 0.7, 1.7, 2.9, 1.8, 'Patient Management', 'Registration, records, and role-based access.', BLUE)
add_box(slide, 3.9, 1.7, 2.9, 1.8, 'AI Triage & Diagnosis', 'Risk scoring, symptom evaluation, and diagnostics.', TEAL)
add_box(slide, 7.1, 1.7, 2.9, 1.8, 'Predictive Maintenance', 'Equipment failure forecasting and device telemetry.', RGBColor(92, 107, 192))
add_box(slide, 10.3, 1.7, 2.9, 1.8, 'Clinical Collaboration', 'Secure chat, attachments, and coordinated care.', RGBColor(33, 150, 243))

# Slide 5: Platform Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'System Architecture')
add_bullets(slide, [
    'Frontend: Next.js interface for dashboards, patient views, and clinical workflows.',
    'Backend: Flask API manages authentication, patient records, diagnostics, and chat.',
    'Database: SQLAlchemy models support users, patients, devices, medical history, and assessments.',
    'Deployment: Docker-ready with support for local use and future cloud scaling.'
], width=6.4)

# Slide 6: AI and Analytics Modules
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'AI & Analytics Modules')
add_box(slide, 0.7, 1.7, 2.9, 2.0, 'Triage Engine', 'Prioritizes patients based on urgency and symptom data.', BLUE)
add_box(slide, 3.9, 1.7, 2.9, 2.0, 'Diagnostic Engine', 'Provides AI-supported analysis for clinical decision support.', TEAL)
add_box(slide, 7.1, 1.7, 2.9, 2.0, 'Equipment Analytics', 'Predicts maintenance risk from telemetry and usage patterns.', RGBColor(76, 175, 80))
add_box(slide, 10.3, 1.7, 2.9, 2.0, 'Admission Forecasting', 'Supports operational planning for likely peak demand periods.', RGBColor(255, 152, 0))

# Slide 7: Impact & Value
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'Value to the Hospital')
add_bullets(slide, [
    'Improves responsiveness by accelerating triage and clinical review.',
    'Reduces downtime and operational risk through predictive maintenance.',
    'Supports better resource planning and patient flow across departments.',
    'Creates a strong foundation for future digital transformation in healthcare.'
], width=6.6)

# Slide 8: Future Roadmap
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'Future Roadmap')
add_bullets(slide, [
    'Expand the AI models using real-world clinical data and feedback.',
    'Integrate with EHRs, imaging systems, and additional hospital platforms.',
    'Improve mobile access, reporting, alerts, and user experience.',
    'Pilot the solution in a real clinical environment for validation and refinement.'
], width=6.7)

# Slide 9: Closing
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'Conclusion', 'A practical and scalable health intelligence solution')
add_bullets(slide, [
    'The platform demonstrates how AI and digital workflows can strengthen healthcare delivery.',
    'It addresses operational efficiency, patient safety, and intelligent decision support.',
    'It is a strong foundation for further academic, clinical, and industrial development.'
], width=7.2)

prs.save(out_file)
print(f'Created: {out_file}')
