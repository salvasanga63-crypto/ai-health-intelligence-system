from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pathlib import Path

root = Path('/home/salva-kil/PROJECTS/AI-driven Health Intelligence System')
out_file = root / 'AI_Health_Intelligence_Pitch_Deck.pptx'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Theme colors
NAVY = RGBColor(12, 41, 68)
BLUE = RGBColor(0, 102, 204)
TEAL = RGBColor(0, 150, 136)
LIGHT = RGBColor(245, 248, 252)
DARK = RGBColor(35, 45, 55)
GRAY = RGBColor(100, 112, 124)


def add_title(slide, title, subtitle=None):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = LIGHT

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.8))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.05), Inches(8.5), Inches(0.5))
        sf = sub_box.text_frame
        sf.clear()
        p2 = sf.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(12)
        p2.font.color.rgb = GRAY


def add_bullets(slide, bullets, left=0.8, top=1.6, width=5.8, height=4.8, title=None):
    if title:
        add_title(slide, title)
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(19)
        p.font.color.rgb = DARK
        p.font.name = 'Calibri'
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)


def add_box(slide, x, y, w, h, title, body, fill_color):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    tx = shape.text_frame
    tx.clear()
    p = tx.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    p2 = tx.add_paragraph()
    p2.text = body
    p2.font.size = Pt(12)
    p2.font.color.rgb = RGBColor(255, 255, 255)
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(6)

# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
# background
bg = slide.background
fill = bg.fill
fill.solid()
fill.fore_color.rgb = NAVY

title = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12.0), Inches(1.1))
tf = title.text_frame
tf.clear()
p = tf.paragraphs[0]
p.text = 'AI-Driven Health Intelligence System'
p.font.size = Pt(28)
p.font.bold = True
p.font.color.rgb = RGBColor(255, 255, 255)

subtitle = slide.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(9.6), Inches(1.0))
sf = subtitle.text_frame
sf.clear()
p2 = sf.paragraphs[0]
p2.text = 'An intelligent platform for patient triage, diagnostics, predictive maintenance, and hospital operations.'
p2.font.size = Pt(18)
p2.font.color.rgb = RGBColor(220, 230, 240)

slide.shapes.add_textbox(Inches(0.7), Inches(3.3), Inches(5.6), Inches(1.7)).text_frame.text = 'Prepared for supervisor evaluation'
# Add small callout box
shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(3.1), Inches(3.5), Inches(2.1))
shape.fill.solid(); shape.fill.fore_color.rgb = TEAL; shape.line.color.rgb = TEAL
shape.text_frame.text = 'Problem\nSolution\nImpact\nScalability'
for p in shape.text_frame.paragraphs:
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(255,255,255)

# Slide 2: Problem
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'Problem Statement')
add_bullets(slide, [
    'Hospitals face delays in patient triage and decision-making.',
    'Limited visibility into equipment health can disrupt care delivery.',
    'Manual record handling makes it hard to track patient risk and care flow.',
    'There is a need for a scalable digital platform that supports clinicians and administrators.'
], width=7.2)

# Slide 3: Solution
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'Proposed Solution')
add_bullets(slide, [
    'A web-based health intelligence platform that combines AI and operational data.',
    'Supports patient management, diagnostics, prediction, and role-based dashboards.',
    'Provides secure access for admins, doctors, nurses, and patients.',
    'Integrates a backend API with a modern frontend for practical hospital use.'
], width=7.0)

# Slide 4: Key Features
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'Key Features')
add_box(slide, 0.7, 1.6, 2.9, 1.8, 'Patient & Staff Management', 'Registration, secure login, role-based access, and patient records.', BLUE)
add_box(slide, 3.8, 1.6, 2.9, 1.8, 'AI Triage & Diagnosis', 'Symptom analysis, diagnostics, nutrition, and risk scoring.', TEAL)
add_box(slide, 6.9, 1.6, 2.9, 1.8, 'Predictive Analytics', 'Peak admissions, failure prediction, and maintenance insights.', RGBColor(92, 107, 192))
add_box(slide, 9.9, 1.6, 2.9, 1.8, 'Clinical Communication', 'Chat rooms, attachments, and coordinated care workflows.', RGBColor(33, 150, 243))

# Slide 5: Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'System Architecture')
add_bullets(slide, [
    'Frontend: Next.js interface for dashboards and clinical workflows.',
    'Backend: Flask API handling authentication, records, predictions, and chat.',
    'Database: SQLAlchemy models for patients, users, devices, and medical records.',
    'Deployment: Docker-ready with support for local development and cloud extension.'
], width=6.2)

# Slide 6: AI/ML Modules
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'AI & Analytics Modules')
add_box(slide, 0.7, 1.6, 2.9, 2.0, 'Triage Engine', 'Prioritizes patients based on risk and symptoms.', RGBColor(0, 102, 204))
add_box(slide, 3.9, 1.6, 2.9, 2.0, 'Diagnostic Engine', 'Evaluates symptoms and health context for decision support.', RGBColor(0, 150, 136))
add_box(slide, 7.1, 1.6, 2.9, 2.0, 'Equipment Maintenance', 'Predicts device failures and operational risk.', RGBColor(76, 175, 80))
add_box(slide, 10.3, 1.6, 2.9, 2.0, 'Admission Forecasting', 'Suggests likely peak demand periods for planning.', RGBColor(255, 152, 0))

# Slide 7: Value & Impact
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'Value to the Hospital')
add_bullets(slide, [
    'Improves triage speed and clinical responsiveness.',
    'Reduces equipment downtime through predictive maintenance.',
    'Supports better resource planning and patient flow.',
    'Creates a strong foundation for future digital health expansion.'
], width=6.5)

# Slide 8: Future Roadmap
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, 'Next Steps')
add_bullets(slide, [
    'Enhance the AI models with real hospital data and feedback.',
    'Add more integrations with EHR and medical imaging systems.',
    'Improve reporting, alerts, and mobile access.',
    'Pilot the platform in a real clinical environment for validation.'
], width=6.4)

prs.save(out_file)
print(f'Created: {out_file}')
